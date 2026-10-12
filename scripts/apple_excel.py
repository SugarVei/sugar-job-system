"""Sugar 求职记录 Apple 简约版。仅依赖 openpyxl；原文件永不覆盖。

把脚本与 Sugar求职记录_2026-09-30.xlsx 放在同一目录，运行：
    python apple_excel.py
也支持 python apple_excel.py "原文件.xlsx" --output "新文件.xlsx"
统计均来自投递明细；待跟进 = 在线测评 + AI面 + HR面 + 待跟进。
Excel 不支持单元格圆角，状态以淡色底纹模拟胶囊视觉。
"""
from __future__ import annotations

import argparse
from collections import Counter
from copy import copy
from datetime import date, datetime
from io import BytesIO
import math
from pathlib import Path
import re
import tempfile
import unicodedata
from zipfile import ZipFile, ZIP_DEFLATED, BadZipFile
import xml.etree.ElementTree as ET

from openpyxl import load_workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import Paragraph, ParagraphProperties, CharacterProperties
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

FONT = '微软雅黑'
TEXT, SECONDARY, HEADER, LINE = '1D1D1F', '86868B', 'F2F3F7', 'E5E5EA'
STATUS = {
    '已投递': ('F2F3F5', '5F6368'), '在线测评': ('E8F4FD', '1973E8'),
    'AI面': ('EEF2FF', '4F46E5'), '一面': ('EEF2FF', '4F46E5'),
    '二面': ('EEF2FF', '4F46E5'), 'HR面': ('E6F4EA', '1E7E34'),
    '已拒绝': ('FDECEE', 'C93B47'), '已放弃': ('FDECEE', 'C93B47'),
    '待投递': ('FEF7EB', 'B25E09'), 'Offer': ('F0F7F2', '2E7D47'),
    '待跟进': ('FAF4EF', 'B25E28'), '人才库': ('F2F3F5', '5F6368'),
}
ORDER = ['待投递', '已投递', '在线测评', 'AI面', 'HR面', '一面', '二面',
         'Offer', '已拒绝', '已放弃', '人才库', '待跟进']
CHART_NS = 'http://schemas.openxmlformats.org/drawingml/2006/chart'
DRAW_NS = 'http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing'
ET.register_namespace('c', CHART_NS)
ET.register_namespace('xdr', DRAW_NS)
ET.register_namespace('a', 'http://schemas.openxmlformats.org/drawingml/2006/main')
ET.register_namespace('r', 'http://schemas.openxmlformats.org/officeDocument/2006/relationships')


def clean(value):
    """None/NaN 安全处理；不改变原始文本、日期、数值或公式。"""
    return None if isinstance(value, float) and not math.isfinite(value) else value


def text(value):
    value = clean(value)
    return '' if value is None else str(value)


def repair_package(path):
    """只在内存修正旧导出器已知的图表 XML 错误，不改动数据 XML。"""
    output = BytesIO()
    with ZipFile(path) as source, ZipFile(output, 'w', ZIP_DEFLATED) as dest:
        bad = source.testzip()
        if bad:
            raise ValueError(f'文件压缩包损坏：{bad}；请重新导出原文件。')
        for entry in source.infolist():
            data = source.read(entry.filename)
            changed = False
            if re.fullmatch(r'xl/drawings/drawing\d+\.xml', entry.filename):
                root = ET.fromstring(data)
                for anchor in root.findall(f'{{{DRAW_NS}}}twoCellAnchor'):
                    for attr in ('moveWithCells', 'sizeWithCells'):
                        if attr in anchor.attrib:
                            del anchor.attrib[attr]
                            changed = True
                    if changed:
                        anchor.set('editAs', 'oneCell')
            elif re.fullmatch(r'xl/charts/chart\d+\.xml', entry.filename):
                root = ET.fromstring(data)
                for axis in root.iter():
                    if axis.tag in {f'{{{CHART_NS}}}catAx', f'{{{CHART_NS}}}valAx'}:
                        for child in axis:
                            if child.tag == f'{{{CHART_NS}}}tickMark':
                                child.tag = f'{{{CHART_NS}}}majorTickMark'
                                changed = True
                plot = root.find(f'.//{{{CHART_NS}}}plotArea')
                if plot is not None:
                    axes = [node.find(f'{{{CHART_NS}}}axId') for node in plot
                            if node.tag in {f'{{{CHART_NS}}}catAx', f'{{{CHART_NS}}}valAx'}]
                    for chart in plot.findall(f'{{{CHART_NS}}}lineChart'):
                        if not chart.findall(f'{{{CHART_NS}}}axId'):
                            for axis_id in axes:
                                if axis_id is not None:
                                    chart.append(copy(axis_id))
                                    changed = True
            if changed:
                data = ET.tostring(root, encoding='utf-8', xml_declaration=True)
            dest.writestr(entry, data)
    output.seek(0)
    return output


def paint(cell, bg='FFFFFF', fg=TEXT, size=10, bold=False, line=True):
    cell.font = Font(name=FONT, size=size, bold=bold, color=fg)
    cell.fill = PatternFill('solid', fgColor=bg)
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell.border = Border(bottom=Side(style='thin', color=LINE)) if line else Border()


def merge_row(ws, row, left, right, value, **style):
    for col in range(left, right + 1):
        paint(ws.cell(row, col), **style)
    ws.merge_cells(start_row=row, start_column=left, end_row=row, end_column=right)
    cell = ws.cell(row, left)
    cell.value = value
    paint(cell, **style)
    return cell


def width(value):
    return max([0] + [sum(2 if unicodedata.east_asian_width(ch) in 'WF' else 1
                           for ch in line) for line in text(value).splitlines()])


def header(ws, row, first, labels):
    for i, label in enumerate(labels):
        cell = ws.cell(row, first + i, label)
        paint(cell, HEADER, '2C2C2E', 10.5, True)
        cell.border = Border(bottom=Side(style='thin', color='D1D1D6'))
    ws.row_dimensions[row].height = 30


def distribution(ws, row, col, labels, pairs, total):
    header(ws, row, col, labels)
    for offset, (label, count) in enumerate(pairs, 1):
        r = row + offset
        for i, value in enumerate((label, count, count / total if total else 0)):
            cell = ws.cell(r, col + i, value)
            paint(cell, 'FBFBFC' if offset % 2 else 'FFFFFF')
            if i == 2:
                cell.number_format = '0.0%'
            if i == 0 and labels[0] == '状态' and label in STATUS:
                paint(cell, *STATUS[label], bold=True)
        ws.row_dimensions[r].height = 24


def key_counts(records, index):
    return Counter(text(row[index]).strip() or '未填写' for row in records)


def sort_counts(counts):
    return sorted(counts.items(), key=lambda pair: (-pair[1], pair[0]))


def build_overview(wb, records, as_of):
    old = wb['投递总览'] if '投递总览' in wb else None
    index = wb.index(old) if old else 0
    if old:
        wb.remove(old)
    ws = wb.create_sheet('投递总览', index)
    statuses, cities, channels = (key_counts(records, i) for i in (5, 2, 3))
    total = len(records)
    interviews = sum(statuses[s] for s in ('AI面', 'HR面', '一面', '二面', 'Offer'))
    offers = statuses['Offer']
    follow = sum(statuses[s] for s in ('在线测评', 'AI面', 'HR面', '待跟进'))
    rate = lambda n: f'{n / total * 100 if total else 0:.1f}%'
    merge_row(ws, 1, 1, 6, 'Sugar 求职系统 · 投递情况总览', size=16, bold=True, line=False)
    merge_row(ws, 2, 1, 6, f'求职进度全景监控看板 · 数据更新至 {as_of}', fg=SECONDARY, size=9.5, line=False)
    ws.row_dimensions[1].height = 36
    ws.row_dimensions[2].height = 18
    cards = [
        ('投递总数', total, '全部记录', 'F0F4FA', '1D6FB8', 'DCE5F2'),
        ('进入面试', interviews, f'进面率 {rate(interviews)}', 'F4F1FA', '6351A2', 'E3DCF2'),
        ('获得 Offer', offers, f'Offer 率 {rate(offers)}' + (' · 冲刺中' if not offers else ''), 'F0F7F2', '2E7D47', 'D9EBDC'),
        ('待跟进事项', follow, '测评 / AI面 / HR面 / 待跟进', 'FAF4EF', 'B25E28', 'F2E4D6'),
    ]
    for i, (label, value, desc, bg, fg, edge) in enumerate(cards):
        row, col = (4 if i < 2 else 8), (1 if i % 2 == 0 else 4)
        for offset, (content, size, bold, height) in enumerate(((label, 10.5, True, 22), (value, 26, True, 36), (desc, 9.5, False, 20))):
            cell = merge_row(ws, row + offset, col, col + 2, content, bg=bg, fg=fg, size=size, bold=bold, line=False)
            if offset == 2:
                cell.border = Border(bottom=Side(style='thin', color=edge))
                for c in range(col + 1, col + 3):
                    ws.cell(row + offset, c).border = copy(cell.border)
            ws.row_dimensions[row + offset].height = height
    # Keep observed statuses, so the real seven statuses fit Row 13:19 without discarding records.
    status_rows = [(s, statuses[s]) for s in ORDER if statuses[s]]
    status_rows += [(s, n) for s, n in sort_counts(statuses) if s not in ORDER]
    distribution(ws, 12, 1, ['状态', '数量', '占比'], status_rows, total)
    distribution(ws, 12, 4, ['城市', '数量', '占比'], sort_counts(cities)[:10], total)
    channel_row = max(24, 14 + len(status_rows))
    channel_rows = sort_counts(channels)
    distribution(ws, channel_row, 1, ['渠道', '数量', '占比'], channel_rows, total)
    end = max(36, channel_row + len(channel_rows), 25)
    shenzhen = sum(n for city, n in cities.items() if '深圳' in city)
    focus = f'以深圳与珠三角为核心\n深圳相关记录 {shenzhen} 条' if shenzhen else '聚焦目标城市与岗位，提高投递匹配度'
    insight = f'求职洞察摘要\n\n① {focus}\n\n② 进面率 {rate(interviews)}，持续复盘面试\n\n③ 建议增加精准内推\n优先跟进 {follow} 项待办'
    for row in ws.iter_rows(min_row=25, max_row=end, min_col=4, max_col=6):
        for cell in row:
            paint(cell, 'F8F9FA', line=False)
    ws.merge_cells(start_row=25, start_column=4, end_row=end, end_column=6)
    ws['D25'] = insight
    paint(ws['D25'], 'F8F9FA', size=11, line=False)
    for col, value in zip('ABCDEF', (24, 12, 16, 24, 12, 16)):
        ws.column_dimensions[col].width = value
    for row in (3, 7, 11, 23):
        if not any(ws.cell(row, c).value is not None for c in range(1, 7)):
            ws.row_dimensions[row].height = 12
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'portrait'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.print_options.horizontalCentered = True
    ws.print_area = f'A1:F{end}'
    return total, interviews, offers, follow


def style_detail(ws):
    for merged in list(ws.merged_cells.ranges):
        # The source title is the only permitted merged area in a detail list.
        if merged.min_row > 1:
            raise ValueError('投递明细中存在合并数据单元格，请先检查数据，避免遗漏。')
        ws.unmerge_cells(str(merged))
    merge_row(ws, 1, 1, 15, text(ws['A1'].value) or 'Sugar 求职系统 · 投递明细', size=16, bold=True, line=False)
    ws.row_dimensions[1].height = 36
    labels = [ws.cell(2, c).value for c in range(1, 16)]
    header(ws, 2, 1, labels)
    widths = [20, 22, 12, 16, 14, 14, 16, 12, 24, 21, 21, 12, 32, 40, 14]
    for row in ws.iter_rows(min_row=3, max_col=15):
        for i, cell in enumerate(row):
            widths[i] = max(widths[i], min(56 if i >= 12 else 40, width(cell.value) + 4))
    for i, value in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = value
    for row in ws.iter_rows(min_row=3, max_col=15):
        lines = 1
        for i, cell in enumerate(row):
            if isinstance(cell.value, float) and not math.isfinite(cell.value):
                cell.value = None
            paint(cell, 'FBFBFC' if cell.row % 2 else 'FFFFFF', bold=i < 2)
            if i == 5 and text(cell.value) in STATUS:
                paint(cell, *STATUS[text(cell.value)], bold=True)
            if isinstance(cell.value, (datetime, date)):
                cell.number_format = 'yyyy-mm-dd hh:mm' if isinstance(cell.value, datetime) and cell.value.time() != datetime.min.time() else 'yyyy-mm-dd'
            lines = max(lines, math.ceil(width(cell.value) / (widths[i] - 2)), len(text(cell.value).splitlines()))
        # 25pt for ordinary rows; expand long JD/notes so all text remains visible.
        ws.row_dimensions[row[0].row].height = min(409, max(25, lines * 15 + 8))
    ws.freeze_panes = 'A3'
    ws.auto_filter.ref = f'A2:O{ws.max_row}'
    ws.print_title_rows = '1:2'
    ws.print_area = f'A1:O{ws.max_row}'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0


def style_weekly(ws):
    for row in ws:
        for cell in row:
            if cell.__class__.__name__ != 'MergedCell':
                paint(cell, line=cell.row >= 3 and cell.column <= 2)
    ws['A1'] = '每周投递趋势'
    paint(ws['A1'], size=16, bold=True, line=False)
    ws.row_dimensions[1].height = 36
    header_row = next((r for r in range(2, min(10, ws.max_row) + 1)
                       if ws.cell(r, 1).value == '周次'), 3)
    for c in (1, 2):
        paint(ws.cell(header_row, c), HEADER, '2C2C2E', 10.5, True)
    for r in range(2, ws.max_row + 1):
        ws.row_dimensions[r].height = 24
        if r > header_row:
            for c in (1, 2):
                paint(ws.cell(r, c), 'FBFBFC' if (r - header_row) % 2 else 'FFFFFF')
    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['B'].width = 16
    ws.column_dimensions['C'].width = 8
    # Regenerate the old exporter's damaged weekly chart with standard openpyxl OOXML.
    ws._charts.clear()
    if ws.max_row > header_row:
        chart = LineChart()
        chart.title = '周度投递趋势'
        chart.style = None
        chart.roundedCorners = False
        chart.graphical_properties = GraphicalProperties(solidFill='FFFFFF', ln=LineProperties(noFill=True))
        chart.x_axis.axPos = 'b'
        chart.y_axis.axPos = 'l'
        for axis in (chart.x_axis, chart.y_axis):
            axis.delete = False
            axis.tickLblPos = 'nextTo'
            axis.majorTickMark = 'none'
            axis.minorTickMark = 'none'
            axis.crosses = 'autoZero'
            axis.spPr = GraphicalProperties(ln=LineProperties(solidFill=LINE, w=9525))
            axis.txPr = RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=CharacterProperties(sz=900, solidFill=SECONDARY)))])
        chart.y_axis.numFmt = '0'
        chart.y_axis.scaling.min = 0
        chart.y_axis.majorGridlines = ChartLines(spPr=GraphicalProperties(ln=LineProperties(solidFill=LINE, w=9525)))
        chart.width, chart.height = 23, 12
        chart.legend = None
        chart.add_data(Reference(ws, min_col=2, min_row=header_row, max_row=ws.max_row), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=header_row + 1, max_row=ws.max_row))
        chart.series[0].graphicalProperties.line.solidFill = '1973E8'
        chart.series[0].graphicalProperties.line.width = 25400
        chart.series[0].marker.symbol = 'circle'
        chart.series[0].marker.size = 5
        chart.series[0].marker.graphicalProperties.solidFill = '1973E8'
        chart.series[0].marker.graphicalProperties.line.solidFill = '1973E8'
        chart.series[0].smooth = False
        ws.add_chart(chart, 'D3')
    ws.freeze_panes = f'A{header_row + 1}'


def snapshot(ws):
    return [(cell.coordinate, clean(cell.value), cell.data_type if clean(cell.value) is not None else 'blank',
             cell.hyperlink.target if cell.hyperlink else None)
            for row in ws.iter_rows(min_row=2, max_col=15) for cell in row]


def transform(source, destination):
    source, destination = Path(source).resolve(), Path(destination).resolve()
    if source == destination:
        raise ValueError('输出路径不能与原文件相同。')
    if not source.is_file():
        raise FileNotFoundError(f'未找到原文件：{source}。请把脚本与原文件放在同一目录，或传入文件路径。')
    wb = load_workbook(repair_package(source), data_only=False, keep_links=True)
    if '投递明细' not in wb:
        raise ValueError('原文件缺少“投递明细”工作表。')
    detail = wb['投递明细']
    if detail.max_column != 15 or detail.cell(2, 6).value != '当前状态':
        raise ValueError('原文件应为表头在第 2 行的 15 列 Sugar 投递明细。')
    before = snapshot(detail)
    records = [tuple(clean(c.value) for c in row) for row in detail.iter_rows(min_row=3, max_col=15)
               if any(clean(c.value) is not None for c in row)]
    if any(cell.data_type == 'f' for row in detail.iter_rows(min_row=3, max_col=15) for cell in row):
        raise ValueError('投递明细含公式；为避免用公式文本统计，请先核对源文件。')
    match = re.search(r'\d{4}-\d{2}-\d{2}', source.name)
    as_of = match.group(0) if match else date.today().isoformat()
    metrics = build_overview(wb, records, as_of)
    style_detail(detail)
    if '周度趋势' in wb:
        style_weekly(wb['周度趋势'])
    for ws in wb:
        for view in ws.views.sheetView:
            view.showGridLines = False
        ws.sheet_properties.outlinePr.summaryRight = False
        if ws.title not in ('投递总览', '投递明细', '周度趋势'):
            for row in ws:
                for cell in row:
                    if cell.__class__.__name__ != 'MergedCell':
                        paint(cell)
        ws.sheet_view.zoomScale = 90 if ws.title == '投递明细' else 100
    wb.active = wb.index(wb['投递总览'])
    if snapshot(detail) != before:
        raise ValueError('明细完整性检查失败；没有保存输出。')
    destination.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(dir=destination.parent, suffix='.xlsx', delete=False) as temp:
        temporary = Path(temp.name)
    try:
        wb.save(temporary)
        with ZipFile(temporary) as zip_file:
            if zip_file.testzip():
                raise ValueError('输出压缩包校验失败。')
        check = load_workbook(temporary, data_only=False)
        if snapshot(check['投递明细']) != before:
            raise ValueError('保存后的明细完整性检查失败。')
        check.close()
        temporary.replace(destination)
    finally:
        temporary.unlink(missing_ok=True)
        wb.close()
    return destination, metrics


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('source', nargs='?', type=Path, default=Path(__file__).resolve().parent / 'Sugar求职记录_2026-09-30.xlsx')
    parser.add_argument('--output', type=Path)
    args = parser.parse_args()
    output = args.output or args.source.resolve().parent / 'Sugar求职记录_Apple简约版.xlsx'
    try:
        path, (total, interviews, offers, follow) = transform(args.source, output)
    except (OSError, ValueError, KeyError, ET.ParseError, BadZipFile, InvalidFileException) as error:
        parser.exit(1, f'处理失败：{error}\n')
    print(f'已生成：{path}\n明细：{total} 条 × 15 列；进入面试 {interviews}；Offer {offers}；待跟进 {follow}。')


if __name__ == '__main__':
    main()
