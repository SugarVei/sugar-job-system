"""Run with python scripts/verify-apple-excel.py (openpyxl + standard library)."""
from datetime import date
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZipFile
import warnings
import xml.etree.ElementTree as ET

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from apple_excel import transform, clean, DRAW_NS

HEADERS = ['公司名称', '岗位名称', '城市', '投递渠道', '投递日期', '当前状态',
           '薪资范围', '优先级', '下一步动作', '下一步时间', '截止时间', '匹配度',
           'JD关键词', '备注', '录入时间']


def source(path, records):
    wb = Workbook()
    wb.active.title = '投递总览'
    detail = wb.create_sheet('投递明细')
    detail.append(['Sugar 求职系统 · 投递明细'])
    detail.append(HEADERS)
    for name, status, applied in records:
        detail.append([name, '岗位', None, '官网', applied, status, 0, 'high', '移除', None, None, 0, '制造、采购', '<>&换行\n备注', date(2026, 9, 30)])
    if records:
        detail['A3'].hyperlink = 'https://example.com/job'
        detail['N3'].comment = Comment('保留原始备注', 'User')
    weekly = wb.create_sheet('周度趋势')
    weekly.append(['每周投递趋势'])
    weekly.append([])
    weekly.append(['周次', '投递数'])
    wb.save(path)


with TemporaryDirectory(prefix='sugar-progress-test-') as task_folder:
    root = Path(task_folder)
    records = [('已拒绝', '已拒绝', '2026-10-12'), ('新 A', '已投递', '2026-10-12'),
               ('HR', 'HR面', '2026-10-12'), ('二面', '二面', '2026-01-01'),
               ('Offer', 'Offer', None), ('旧', '已投递', '2026-09-30'),
               ('新 B', '已投递', '2026-10-12'), ('空日期', '已投递', None)]
    src, dest = root / 'source.xlsx', root / 'output.xlsx'
    source(src, records)
    _, metrics = transform(src, dest)
    assert metrics == (8, 3, 1, 1)
    with warnings.catch_warnings():
        warnings.filterwarnings('ignore', message='DrawingML support is incomplete.*')
        wb = load_workbook(dest)
    detail = wb['投递明细']
    expected = ['Offer', '二面', 'HR', '新 A', '新 B', '旧', '空日期', '已拒绝']
    assert [detail.cell(r, 2).value for r in range(3, 11)] == expected
    assert detail.max_column == 11
    assert detail.freeze_panes == 'B3'
    assert detail.auto_filter.ref == 'B2:K10'
    assert [detail.cell(2, c).value for c in range(2, 12)] == [HEADERS[i] for i in (0,1,2,3,4,5,6,12,13,14)]
    assert detail['B10'].hyperlink.target == 'https://example.com/job'
    assert detail['B10'].hyperlink.ref == 'B10'
    assert detail['J10'].comment.text == '保留原始备注'
    assert detail['H3'].value == 0
    assert detail['J3'].value == '<>&换行\n备注'
    assert all(detail.cell(r, 1).value is None for r in range(1, 11))
    assert all(sheet.sheet_view.showGridLines is False for sheet in wb)
    with ZipFile(dest) as package:
        drawing = ET.fromstring(package.read('xl/drawings/statusProgress.xml'))
        assert len(drawing.findall(f'.//{{{DRAW_NS}}}grpSp')) == 7
        ids = [node.attrib['id'] for node in drawing.iter(f'{{{DRAW_NS}}}cNvPr')]
        assert len(ids) == len(set(ids))
    wb.close()
    source(src, [])
    _, metrics = transform(src, dest)
    assert metrics == (0, 0, 0, 0)
    with ZipFile(dest) as package:
        assert 'xl/drawings/statusProgress.xml' not in package.namelist()
    assert clean(float('nan')) is None
    assert clean(float('inf')) is None
print('PASS: Python progress/date/stable sorting, ten-field projection, links and comments, numeric zero, empty and NaN handling.')
